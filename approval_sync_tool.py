import argparse
import json
import os
import re
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DEFAULT_LEVELS = "Level 1:level1Approvers,Level 2:level2Approvers,Level 3:level3Approvers,Level 4:level4Approvers,Level 5:level5Approvers"


def norm(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).lower()


def parse_levels(value):
    result = {}
    for pair in value.split(","):
        if not pair.strip():
            continue
        column, field = pair.split(":", 1)
        result[column.strip()] = field.strip()
    return result


def read_config(path):
    with open(path, "r", encoding="utf-8") as f:
        raw = json.load(f)
    aliases = {}
    for row in raw.get("aliases", []):
        source = norm(row.get("sourceApprover"))
        if source:
            aliases[source] = row.get("aliases", [])
    return {"skipApprovers": [norm(x) for x in raw.get("skipApprovers", [])], "aliases": aliases}


def header_index(ws, row):
    return {str(cell.value).strip(): idx + 1 for idx, cell in enumerate(ws[row]) if cell.value is not None}


def find_sheet(wb, preferred_name):
    if preferred_name in wb.sheetnames:
        return wb[preferred_name]
    lowered = preferred_name.lower()
    for sheet_name in wb.sheetnames:
        if lowered in sheet_name.lower():
            return wb[sheet_name]
    raise ValueError(f"Cannot find sheet: {preferred_name}")


def source_name_index(ws, headers, employee_name_col, preferred_name_col):
    by_name = {}
    for row in range(3, ws.max_row + 1):
        person = {
            "sourceEmployeeRow": row,
            "employeeName": ws.cell(row, headers[employee_name_col]).value,
            "preferredName": ws.cell(row, headers[preferred_name_col]).value if preferred_name_col in headers else None,
        }
        for field in ("employeeName", "preferredName"):
            key = norm(person[field])
            if key:
                by_name.setdefault(key, []).append(person)
    return by_name


def pick_aliases(source_approver, config, names):
    source_key = norm(source_approver)
    if source_key in config["aliases"]:
        return config["aliases"][source_key], "manual alias config"

    hits = names.get(source_key, [])
    unique = {}
    for hit in hits:
        unique[(norm(hit["employeeName"]), norm(hit["preferredName"]))] = hit
    hits = list(unique.values())
    if len(hits) == 1:
        hit = hits[0]
        candidates = [hit.get("preferredName"), hit.get("employeeName"), source_approver]
        return list(dict.fromkeys([str(x).strip() for x in candidates if norm(x)])), "employee/preferred name crosscheck"

    return [str(source_approver).strip()], "source approver only"


def should_skip(value, skip_terms):
    text = norm(value)
    return any(term and term in text for term in skip_terms)


def build_plan(args):
    levels = parse_levels(args.levels)
    config = read_config(args.config)
    wb = openpyxl.load_workbook(args.source, data_only=True)
    ws = find_sheet(wb, args.sheet)
    headers = header_index(ws, args.header_row)

    required = [args.employee_id_col, args.employee_name_col] + list(levels.keys())
    missing = [col for col in required if col not in headers]
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(missing))

    names = source_name_index(ws, headers, args.employee_name_col, args.preferred_name_col)
    plan, skipped, review, seen = [], [], [], set()

    for row in range(args.header_row + 1, ws.max_row + 1):
        employee_id = ws.cell(row, headers[args.employee_id_col]).value
        employee_name = ws.cell(row, headers[args.preferred_name_col]).value if args.preferred_name_col in headers else None
        employee_name = employee_name or ws.cell(row, headers[args.employee_name_col]).value

        if not employee_id:
            review.append({"sourceRow": row, "employeeName": employee_name, "reason": "missing employee id"})
            continue

        for source_col, target_field in levels.items():
            source_approver = ws.cell(row, headers[source_col]).value
            if not norm(source_approver):
                continue

            unique_key = (str(employee_id), target_field, norm(source_approver))
            if unique_key in seen:
                continue
            seen.add(unique_key)

            base = {
                "sourceRow": row,
                "employeeId": str(employee_id),
                "employeeName": str(employee_name or "").strip(),
                "sourceColumn": source_col,
                "field": target_field,
                "sourceApprover": str(source_approver).strip(),
            }

            if should_skip(source_approver, config["skipApprovers"]):
                skipped.append({**base, "reason": "approver is in skip list"})
                continue

            aliases, basis = pick_aliases(source_approver, config, names)
            plan.append({**base, "aliases": aliases, "basis": basis})

    return plan, skipped, review


def add_table_sheet(wb, name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])

    fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="D9D9D9")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cells in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in cells:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center")
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = min(max(len(str(headers[col - 1])) + 8, 14), 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_review(path, plan, skipped, review):
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Summary"
    for row in [["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")], ["Planned rows", len(plan)], ["Skipped rows", len(skipped)], ["Review rows", len(review)]]:
        summary.append(row)
    add_table_sheet(wb, "Plan", ["sourceRow", "employeeId", "employeeName", "sourceColumn", "field", "sourceApprover", "aliases", "basis"], [{**row, "aliases": " | ".join(row.get("aliases", []))} for row in plan])
    add_table_sheet(wb, "Skipped", ["sourceRow", "employeeId", "employeeName", "sourceColumn", "field", "sourceApprover", "reason"], skipped)
    add_table_sheet(wb, "Review", ["sourceRow", "employeeName", "reason"], review)
    wb.save(path)


def write_runner(path, template_path, plan):
    with open(template_path, "r", encoding="utf-8") as f:
        template = f.read()
    with open(path, "w", encoding="utf-8") as f:
        f.write("window.__APPROVAL_SYNC_PLAN__ = ")
        f.write(json.dumps(plan, ensure_ascii=False, indent=2))
        f.write(";\n")
        f.write(template)


def main():
    parser = argparse.ArgumentParser(description="Prepare approval or hierarchy update plans from Excel.")
    parser.add_argument("--source", required=True, help="Source Excel workbook path")
    parser.add_argument("--sheet", required=True, help="Source sheet name")
    parser.add_argument("--header-row", type=int, default=2, help="Header row number")
    parser.add_argument("--employee-id-col", default="Employee ID", help="Employee id column name")
    parser.add_argument("--employee-name-col", default="Employee Name", help="Employee name column name")
    parser.add_argument("--preferred-name-col", default="Preferred Name", help="Preferred name column name")
    parser.add_argument("--levels", default=DEFAULT_LEVELS, help="Comma-separated mapping: source column:target field")
    parser.add_argument("--config", default="alias_config.json", help="Alias config path")
    parser.add_argument("--output-dir", default="output", help="Output folder")
    args = parser.parse_args()

    script_dir = os.path.dirname(os.path.abspath(__file__))
    if not os.path.isabs(args.config):
        args.config = os.path.join(script_dir, args.config)
    output_dir = args.output_dir if os.path.isabs(args.output_dir) else os.path.join(script_dir, args.output_dir)
    os.makedirs(output_dir, exist_ok=True)

    plan, skipped, review = build_plan(args)
    plan_path = os.path.join(output_dir, "approval_sync_plan.json")
    review_path = os.path.join(output_dir, "approval_sync_review.xlsx")
    runner_path = os.path.join(output_dir, "browser_runner_with_plan.js")

    with open(plan_path, "w", encoding="utf-8") as f:
        json.dump(plan, f, ensure_ascii=False, indent=2)
    write_review(review_path, plan, skipped, review)
    write_runner(runner_path, os.path.join(script_dir, "browser_apply_runner.js"), plan)

    print("Files generated:")
    print("Plan:", plan_path)
    print("Review workbook:", review_path)
    print("Browser runner:", runner_path)
    print({"planned": len(plan), "skipped": len(skipped), "review": len(review)})


if __name__ == "__main__":
    main()
